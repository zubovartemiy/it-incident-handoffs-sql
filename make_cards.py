"""Builds the one-page project card (EN and CZ) as Word documents and exports them to PDF via Word."""
import os
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ORANGE = RGBColor(0xE8, 0x89, 0x2B)

TEXT = {
    "en": {
        "title": "What does a handoff cost a service desk?",
        "sub": "SQL analysis of a public IT incident log: 141 712 events, 24 918 incidents (ServiceNow export, UCI no. 498, CC BY 4.0). Personal project, October 2026.",
        "q": ("Question", "How are handoffs between support groups related to resolution time and SLA breaches, and is ping-pong (a return to an earlier group) worse than a handoff to a new group?"),
        "m": ("What I did", "Typed the raw event log and rebuilt, in SQL with window functions and CTEs, each incident's path through support groups. Compared medians and SLA breaches by number of handoffs (priority 3 only), compared ping-pong with forward-only paths at the same number of handoffs, wrote 11 SQL checks of incident-management process logic and data quality, and a reusable service desk KPI pack. Verified the headline numbers with an independent recount from the raw file."),
        "r": ("Results", [
            "Solved by the first group: median 1 hour, SLA breached in 26 %. With 4 or more handoffs: 282 hours and 81 % (priority 3, 21 644 incidents with a recorded group).",
            "Ping-pong looks worse mainly because bounced incidents have more handoffs. At the same count the differences become small and are not consistent (2 handoffs: 60 % vs 60 % breached).",
            "The tool enforces its own logic (priority matrix, timestamps, SLA flag are clean), but 6.2 % of closed incidents carry no resolution time and drop out of reports based on it.",
        ]),
        "w": ("What it means", "First-group resolution goes with the shortest resolution times; getting incidents to the right group sooner looks more useful than forbidding returns; the resolution time should be mandatory at closure."),
        "t": ("Tools", "SQL (SQLite: window functions, CTEs), Python for loading and charts."),
        "file": "Project_card_IT_incidents_EN",
    },
    "cz": {
        "title": "Kolik stojí předání incidentu mezi skupinami?",
        "sub": "Analýza veřejného logu IT incidentů v SQL: 141 712 událostí, 24 918 incidentů (export ze ServiceNow, UCI č. 498, CC BY 4.0). Vlastní projekt, říjen 2026.",
        "q": ("Otázka", "Jak souvisí předání incidentu mezi skupinami podpory s dobou řešení a porušením SLA a je „ping-pong“ (návrat do skupiny, kde už incident byl) horší než předání nové skupině?"),
        "m": ("Postup", "Z logu událostí jsem v SQL (window funkce, CTE) zrekonstruoval cestu každého incidentu přes skupiny podpory. Porovnal jsem mediány a porušení SLA podle počtu předání (jen priorita 3), porovnal ping-pong s cestami bez návratu při stejném počtu předání, zapsal 11 kontrol procesu řízení incidentů a kvality dat v SQL a připravil sadu KPI pro service desk. Hlavní čísla jsem ověřil nezávislým přepočtem z původního souboru."),
        "r": ("Výsledky", [
            "Vyřešeno první skupinou: medián 1 hodina, SLA porušeno u 26 %. Při 4 a více předáních: 282 hodin a 81 % (priorita 3, 21 644 incidentů se zaznamenanou skupinou).",
            "Ping-pong vypadá hůř hlavně proto, že vrácené incidenty mají více předání. Při stejném počtu předání jsou rozdíly malé a nemají stálý směr (2 předání: 60 % proti 60 %).",
            "Systém hlídá vlastní logiku (matice priorit, časy, příznak SLA jsou čisté), ale 6,2 % uzavřených incidentů nemá čas vyřešení a vypadává z reportů, které s ním počítají.",
        ]),
        "w": ("Co z toho plyne", "Vyřešení první skupinou jde ruku v ruce s nejkratší dobou řešení; užitečnější než zakázat vracení se jeví směrovat incident hned správné skupině; čas vyřešení by měl být při uzavření povinný."),
        "t": ("Nástroje", "SQL (SQLite: window funkce, CTE), Python pro načtení dat a grafy."),
        "file": "Projektova_karta_IT_incidenty_CZ",
    },
}


def build(lang):
    t = TEXT[lang]
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(1.6)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    st.paragraph_format.space_after = Pt(3)

    p = d.add_paragraph()
    r = p.add_run(t["title"]); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY
    p = d.add_paragraph()
    r = p.add_run(t["sub"]); r.italic = True; r.font.size = Pt(9.5)

    def block(head, body):
        p = d.add_paragraph(); p.paragraph_format.space_before = Pt(6)
        r = p.add_run(head.upper()); r.bold = True; r.font.color.rgb = ORANGE; r.font.size = Pt(10)
        if isinstance(body, list):
            for b in body:
                d.add_paragraph(b, style="List Bullet")
        else:
            d.add_paragraph(body)

    block(*t["q"])
    block(*t["m"])
    d.add_picture(str(HERE / "figures" / ("01_handoffs_sla_cz.png" if lang == "cz" else "01_handoffs_sla.png")), width=Cm(15.5))
    block(*t["r"])
    block(*t["w"])
    block(*t["t"])
    out = HERE / "report" / (t["file"] + ".docx")
    d.save(out)
    return out


import pythoncom
import win32com.client as win32
pythoncom.CoInitialize()
word = win32.DispatchEx("Word.Application"); word.Visible = False; word.DisplayAlerts = 0
try:
    for lang in ("en", "cz"):
        docx = build(lang)
        doc = word.Documents.Open(os.path.abspath(docx), ReadOnly=True)
        pdf = os.path.abspath(docx.with_suffix(".pdf"))
        doc.ExportAsFixedFormat(OutputFileName=pdf, ExportFormat=17)
        pages = doc.ComputeStatistics(2)
        doc.Close(SaveChanges=0)
        print(lang, "->", pdf, "pages:", pages)
finally:
    word.Quit()
